"""Document text extraction and normalization utilities."""
from __future__ import annotations

import hashlib
import re
import unicodedata
from io import BytesIO
from pathlib import Path


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from a supported TXT, PDF, or DOCX file."""
    extension = Path(filename).suffix.lower()

    if extension == ".txt":
        # UTF-8 is preferred; fallback keeps common legacy text files usable.
        try:
            return content.decode("utf-8-sig")
        except UnicodeDecodeError:
            return content.decode("latin-1", errors="replace")

    if extension == ".pdf":
        import fitz  # PyMuPDF
        chunks = []
        with fitz.open(stream=content, filetype="pdf") as pdf:
            for page in pdf:
                chunks.append(page.get_text("text"))
        return "\n".join(chunks)

    if extension == ".docx":
        from docx import Document
        document = Document(BytesIO(content))
        chunks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(chunks)

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


def normalize_text(text: str) -> str:
    """Normalize Unicode, case, whitespace, and punctuation for comparison."""
    text = unicodedata.normalize("NFKC", text).lower()
    text = re.sub(r"\s+", " ", text)
    # Keep letters/numbers and whitespace; replace punctuation with spaces.
    text = re.sub(r"[^\w\s]", " ", text, flags=re.UNICODE)
    text = re.sub(r"_+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def file_fingerprint(content: bytes) -> str:
    """Return a SHA-256 fingerprint of the original file bytes."""
    return hashlib.sha256(content).hexdigest()
